from __future__ import annotations

import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES = ROOT / "packages" / "gotthard-skills" / "templates"


def build_letter(path: Path, lang: str) -> None:
    d = docx.Document()
    sec = d.sections[0]
    sec.page_height, sec.page_width = Mm(297), Mm(210)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Mm(25), Mm(20), Mm(20), Mm(20)
    style = d.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10.5)

    head = sec.header.paragraphs[0]
    r = head.add_run("{{FIRMA}}")
    r.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    head.add_run("\n{{ABSENDER_ADRESSE}}").font.size = Pt(8.5)
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.add_run("{{FUSSZEILE}}").font.size = Pt(8)

    d.add_paragraph().paragraph_format.space_after = Mm(18)
    d.add_paragraph("{{EMPFAENGER}}")
    p = d.add_paragraph("{{ORT_DATUM}}")
    p.paragraph_format.space_before = Mm(12)
    p.paragraph_format.space_after = Mm(8)
    p = d.add_paragraph()
    p.add_run("{{BETREFF}}").bold = True
    p.paragraph_format.space_after = Mm(6)
    d.add_paragraph("{{ANREDE}}")
    d.add_paragraph("{{INHALT}}")
    p = d.add_paragraph("{{GRUSS}}")
    p.paragraph_format.space_before = Mm(6)
    p = d.add_paragraph("{{SIGNATUR}}")
    p.paragraph_format.space_before = Mm(14)
    d.core_properties.title = "Briefvorlage" if lang == "de" else "Modèle de lettre"
    d.save(path)


def build_templates() -> None:
    TEMPLATES.mkdir(parents=True, exist_ok=True)
    build_letter(TEMPLATES / "Briefvorlage.docx", "de")
    build_letter(TEMPLATES / "Modele_lettre.docx", "fr")
    print("templates written to", TEMPLATES)


SCANS = [
    ("Eingang/scan_0001.pdf", [
        "Labor Risch AG, Bern", "Laborbefund", "Patient: Hans Weber, geb. 03.07.1958",
        "Auftraggeber: Praxis Dr. Meier, Zürich", "Entnahme: 18.09.2026  ·  Befunddatum: 19.09.2026",
        "HbA1c 7.4 %  (Referenz < 5.7)", "LDL-Cholesterin 3.9 mmol/l  (Ziel < 2.6)", "Kreatinin 88 µmol/l  (Referenz 59–104)",
    ]),
    ("Eingang/scan_0002.pdf", [
        "Universitätsspital Zürich – Klinik für Kardiologie", "Austrittsbericht", "Patientin: Claire Dubois, geb. 14.03.1985",
        "An: Dr. med. Anna Meier, Praxis Dr. Meier, Bahnhofstrasse 20, 8001 Zürich", "Hospitalisation: 02.10.2026 – 05.10.2026",
        "Diagnose: Synkope unklarer Ätiologie, Holter-EKG unauffällig", "Procedere: Kontrolle in der Hausarztpraxis in 2 Wochen",
        "Datum: 06.10.2026",
    ]),
    ("Eingang/scan_0003.pdf", [
        "Helsana Versicherungen AG", "Kostengutsprache", "Versicherte Person: Hans Weber, geb. 03.07.1958, Vers.-Nr. 80756123",
        "Leistung: Ernährungsberatung, 6 Sitzungen", "Gültig bis: 31.03.2027", "Datum: 25.09.2026",
    ]),
]

DOSSIER_DUBOIS = [
    ("Dossier patient (fictif)", "Title"),
    ("Patiente : Claire Dubois, née le 14.03.1985. N° AVS 756.1234.5678.97.", "BodyText"),
    ("Médecin traitant : Dr Anna Meier, Zurich.", "BodyText"),
    ("Suivi psychologique depuis le 02.09.2024 (Marc Rochat, psychologue FSP), adressée par Dr Meier.", "BodyText"),
    ("Indication : épisode dépressif moyen avec anxiété, après un licenciement en juin 2024. Troubles du sommeil, fatigue.", "BodyText"),
    ("Thérapie : thérapie cognitivo-comportementale, 1 séance hebdomadaire de 50 minutes.", "BodyText"),
    ("Médication : sertraline 50 mg (prescrite par Dr Meier).", "BodyText"),
    ("Hospitalisation du 02.10.2026 au 05.10.2026 (cardiologie, USZ) pour une syncope, bilan sans particularité.", "BodyText"),
]


def demo(target: Path) -> None:
    target.mkdir(parents=True, exist_ok=True)
    for rel_path, lines in SCANS:
        p = target / rel_path
        p.parent.mkdir(parents=True, exist_ok=True)
        c = canvas.Canvas(str(p), pagesize=A4)
        y = A4[1] - 25 * mm
        for i, line in enumerate(lines):
            c.setFont("Helvetica-Bold" if i < 2 else "Helvetica", 12 if i < 2 else 10)
            c.drawString(22 * mm, y, line)
            y -= 9 * mm
        c.save()

    dubois = target / "Patienten" / "Dubois_Claire"
    weber = target / "Patienten" / "Weber_Hans"
    dubois.mkdir(parents=True, exist_ok=True)
    weber.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    SimpleDocTemplate(str(dubois / "Dossier_Dubois_Claire.pdf"), pagesize=A4).build(
        [Paragraph(t, styles[st]) for t, st in DOSSIER_DUBOIS])

    note = docx.Document()
    note.add_heading("Konsultation 12.09.2026 – Hans Weber", 1)
    note.add_paragraph("Patient: Hans Weber, geb. 03.07.1958, Seestrasse 4, 8810 Horgen. Tel. 044 123 45 67.")
    note.add_paragraph("Grund: Jahreskontrolle Diabetes mellitus Typ 2.")
    note.add_paragraph("Vereinbart: Laborkontrolle (HbA1c, Lipide, Kreatinin), Anmeldung Ernährungsberatung, "
                       "nächste Konsultation am 05.11.2026 um 14:30 bei Dr. Meier.")
    note.add_paragraph("Der Patient wünscht Korrespondenz per Post an die Wohnadresse.")
    note.save(weber / "Konsultation_2026-09-12.docx")

    (target / "Vorlagen").mkdir(exist_ok=True)
    (target / "Vorlagen" / "LIESMICH.txt").write_text(
        "Legen Sie hier Ihre eigenen Vorlagen ab (z. B. Briefvorlage.docx mit dem Briefkopf der Praxis).\n"
        "Gotthard verwendet Vorlagen in diesem Ordner vor den mitgelieferten.\n", encoding="utf-8")
    print("demo workspace written to", target)


if __name__ == "__main__":
    if len(sys.argv) >= 2 and sys.argv[1] == "templates":
        build_templates()
    elif len(sys.argv) >= 3 and sys.argv[1] == "demo":
        demo(Path(sys.argv[2]).expanduser().resolve())
    else:
        raise SystemExit("usage: build_assets.py templates | demo <target-dir>")
