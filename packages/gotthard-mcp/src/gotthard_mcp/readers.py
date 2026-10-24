from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path

import docx
import openpyxl
import pdfplumber
from docx.table import Table
from docx.text.paragraph import Paragraph

TEXT_SUFFIXES = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".eml"}
DOC_SUFFIXES = {".pdf", ".docx", ".xlsx", ".xlsm"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".heic", ".bmp"}
INDEXABLE = TEXT_SUFFIXES | DOC_SUFFIXES


@dataclass
class Segment:
    location: str
    text: str


def extract(path: Path, max_chars: int | None = None) -> list[Segment]:
    s = path.suffix.lower()
    if s == ".pdf":
        segs = _pdf(path)
    elif s == ".docx":
        segs = _docx(path)
    elif s in (".xlsx", ".xlsm"):
        segs = _xlsx(path)
    elif s in (".csv", ".tsv"):
        segs = [Segment("", _csv(path, "\t" if s == ".tsv" else None))]
    elif s in TEXT_SUFFIXES:
        segs = [Segment("", path.read_text(encoding="utf-8", errors="replace"))]
    elif s in IMAGE_SUFFIXES:
        segs = [Segment("", "[Bilddatei – Texterkennung (OCR) ist in dieser Version nicht verfügbar / image file – OCR not available in this version]")]
    else:
        raise ValueError(f"Unsupported file type: {s}")
    if max_chars is not None:
        out, used = [], 0
        for seg in segs:
            if used >= max_chars:
                out.append(Segment("…", f"[truncated after {max_chars} characters]"))
                break
            out.append(Segment(seg.location, seg.text[: max_chars - used]))
            used += len(seg.text)
        segs = out
    return segs


def to_text(segs: list[Segment]) -> str:
    parts = []
    for seg in segs:
        parts.append(f"## {seg.location}\n{seg.text}" if seg.location else seg.text)
    return "\n\n".join(parts).strip()


def _pdf(path: Path) -> list[Segment]:
    segs = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages, 1):
            text = page.extract_text() or ""
            tables = []
            try:
                for t in page.extract_tables():
                    rows = [" | ".join((c or "").replace("\n", " ") for c in row) for row in t if row]
                    if rows:
                        tables.append("\n".join(rows))
            except Exception:
                pass
            body = text
            if tables:
                body += "\n\n[Tabellen / tables]\n" + "\n\n".join(tables)
            if not text.strip() and not tables:
                body = "[Keine Textebene – vermutlich ein Scan / no text layer, probably a scan]"
            segs.append(Segment(f"Seite {i}", body))
    return segs


def _docx(path: Path) -> list[Segment]:
    d = docx.Document(str(path))
    lines: list[str] = []
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            t = p.text.strip()
            if not t:
                continue
            style = (p.style.name or "") if p.style is not None else ""
            if style.lower().startswith(("heading", "überschrift", "titre", "title")):
                lines.append(f"# {t}")
            elif "list" in style.lower():
                lines.append(f"- {t}")
            else:
                lines.append(t)
        elif tag == "tbl":
            tbl = Table(child, d)
            for row in tbl.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
            lines.append("")
    return [Segment("", "\n".join(lines))]


def _xlsx(path: Path) -> list[Segment]:
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    segs = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            if row is None or all(v is None for v in row):
                continue
            rows.append(" | ".join("" if v is None else _fmt(v) for v in row))
        segs.append(Segment(f"Blatt {ws.title}", "\n".join(rows)))
    wb.close()
    return segs


def _fmt(v) -> str:
    if isinstance(v, float):
        return f"{v:.2f}" if abs(v - round(v)) > 1e-9 else str(int(round(v)))
    if hasattr(v, "strftime"):
        return v.strftime("%d.%m.%Y")
    return str(v)


def _csv(path: Path, delimiter: str | None) -> str:
    raw = path.read_text(encoding="utf-8-sig", errors="replace")
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ";" if raw.count(";") > raw.count(",") else ","
    rows = csv.reader(io.StringIO(raw), delimiter=delimiter)
    return "\n".join(" | ".join(r) for r in rows)
