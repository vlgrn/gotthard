from __future__ import annotations

import datetime as dt
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

import docx
import openpyxl
from docx.enum.text import WD_BREAK
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from mcp.server.mcpserver import MCPServer
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import ListFlowable, ListItem, PageBreak, SimpleDocTemplate, Spacer, TableStyle
from reportlab.platypus import Paragraph as PdfParagraph
from reportlab.platypus import Table as PdfTable

from . import readers
from .common import WorkspaceError, log_event, output_path, rel, resolve_in_workspace, versioned_path, workspace_root

server = MCPServer(
    name="gotthard-office",
    instructions=(
        "Office file tools confined to the user's workspace. Paths are relative to the workspace root. "
        "Existing files are never overwritten: outputs get a _v2, _v3 ... suffix unless overwrite=true "
        "(only use overwrite when the user explicitly asked). Always report the returned path to the user."
    ),
)

TEMPLATE_DIRS_ENV = "GOTTHARD_TEMPLATES"


@server.tool()
def read_document(path: str, max_chars: int = 60000) -> str:
    """Read text from a PDF, Word (.docx), Excel (.xlsx), CSV, TXT or Markdown file in the workspace.
    PDFs are returned page by page (incl. detected tables), Excel sheet by sheet.
    Bundled templates can be read by file name, e.g. "Briefvorlage.docx"."""
    try:
        p = resolve_in_workspace(path, must_exist=True)
    except WorkspaceError:
        p = _find_template(path)
    segs = readers.extract(p, max_chars=max_chars)
    return f"# {rel(p)}\n\n{readers.to_text(segs)}"


@server.tool()
def create_xlsx(path: str, sheets: list[dict[str, Any]], overwrite: bool = False) -> dict[str, Any]:
    """Create a formatted Excel workbook.

    sheets: list of {
      "name": "Zusammenfassung",
      "columns": ["Datum", "Text", "Konto", "Betrag"],
      "rows": [["03.01.2026", "Miete", "6000", 2400.0], ...],   # values; strings starting with "=" are formulas
      "number_format": {"D": "#,##0.00"},
      "column_widths": {"B": 40},
      "totals": ["D"]
    }
    Dates as "DD.MM.YYYY" or "YYYY-MM-DD" become real Excel dates. Amounts should be numbers, not text.
    """
    if not sheets:
        raise ValueError("At least one sheet is required")
    out = output_path(path, overwrite, ".xlsx")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    thin = Side(style="thin", color="999999")
    summary = []
    for i, spec in enumerate(sheets):
        name = _sheet_name(spec.get("name") or f"Blatt{i + 1}")
        ws = wb.create_sheet(name)
        columns = spec.get("columns") or []
        rows = spec.get("rows") or []
        start = 1
        if columns:
            ws.append([str(c) for c in columns])
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            start = 2
        for row in rows:
            ws.append([_cell_value(v) for v in (row if isinstance(row, list) else [row])])
        last = ws.max_row
        for row in ws.iter_rows(min_row=start, max_row=last):
            for cell in row:
                if isinstance(cell.value, (dt.date, dt.datetime)):
                    cell.number_format = "DD.MM.YYYY"
                elif isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
        for col, fmt in (spec.get("number_format") or {}).items():
            for cell in ws[col.upper()][start - 1 :]:
                cell.number_format = fmt
        for idx, title in enumerate(columns, 1):
            if not _DATE_HEADER.search(str(title)):
                continue
            for cell in ws[get_column_letter(idx)][start - 1 :]:
                v = cell.value
                if isinstance(v, (int, float)) and not isinstance(v, bool) and 20000 < v < 80000:
                    cell.value = (dt.datetime(1899, 12, 30) + dt.timedelta(days=int(v))).date()
                if isinstance(cell.value, (dt.date, dt.datetime)):
                    cell.number_format = "DD.MM.YYYY"
        totals = [c.upper() for c in (spec.get("totals") or [])]
        if totals and last >= start:
            tr = last + 1
            ws.cell(row=tr, column=1, value="Total").font = Font(bold=True)
            for col in totals:
                c = ws[f"{col}{tr}"]
                c.value = f"=SUM({col}{start}:{col}{last})"
                c.font = Font(bold=True)
                c.number_format = (spec.get("number_format") or {}).get(col, "#,##0.00")
                c.border = Border(top=thin)
        if columns and last >= start:
            ws.auto_filter.ref = f"A1:{get_column_letter(max(1, ws.max_column))}{last}"
        widths = {k.upper(): v for k, v in (spec.get("column_widths") or {}).items()}
        for idx in range(1, ws.max_column + 1):
            letter = get_column_letter(idx)
            if letter in widths:
                ws.column_dimensions[letter].width = widths[letter]
            else:
                longest = max((len(_display(c.value)) for c in ws[letter]), default=8)
                ws.column_dimensions[letter].width = min(max(10, longest + 2), 60)
        summary.append({"sheet": name, "rows": len(rows)})
    wb.save(out)
    log_event("files", action="create", path=rel(out), tool="create_xlsx")
    return {"path": rel(out), "sheets": summary}


def _sheet_name(n: str) -> str:
    return re.sub(r"[\[\]\:\*\?\/\\]", "-", n)[:31]


_DATE_DMY = re.compile(r"^(\d{1,2})\.(\d{1,2})\.(\d{4})$")
_DATE_ISO = re.compile(r"^(\d{4})-(\d{2})-(\d{2})$")
_NUM = re.compile(r"^-?[\d']+(\.\d+)?$")
_DATE_HEADER = re.compile(r"datum|date|data", re.I)


def _cell_value(v: Any) -> Any:
    if isinstance(v, str):
        s = v.strip()
        if s.startswith("="):
            return s
        if m := _DATE_DMY.match(s):
            try:
                return dt.date(int(m[3]), int(m[2]), int(m[1]))
            except ValueError:
                return s
        if m := _DATE_ISO.match(s):
            try:
                return dt.date(int(m[1]), int(m[2]), int(m[3]))
            except ValueError:
                return s
        if _NUM.match(s) and ("'" in s or "." in s):
            try:
                return float(s.replace("'", ""))
            except ValueError:
                return s
        return s
    if isinstance(v, int) and not isinstance(v, bool):
        return v
    if isinstance(v, float):
        return v
    return v


def _display(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:,.2f}"
    if isinstance(v, (dt.date, dt.datetime)):
        return "00.00.0000"
    return str(v)


@server.tool()
def create_docx(
    path: str,
    blocks: list[dict[str, Any]],
    template: str | None = None,
    fields: dict[str, str] | None = None,
    overwrite: bool = False,
) -> dict[str, Any]:
    """Create a Word document, optionally from a template (letterhead).

    blocks: list of
      {"type": "heading", "text": "...", "level": 1}
      {"type": "paragraph", "text": "Text with **bold** parts"}
      {"type": "bullets", "items": ["...", "..."]}
      {"type": "numbered", "items": ["...", "..."]}
      {"type": "table", "rows": [["Kopf", "Kopf"], ["a", "b"]], "header": true}
      {"type": "page_break"}
    template: a .docx in the workspace (e.g. "Vorlagen/Briefvorlage.docx") or a bundled one by name
      ("Briefvorlage.docx", "Modele_lettre.docx"). Placeholders like {{DATUM}} are replaced from
      `fields`; the blocks are inserted at the paragraph containing {{INHALT}} (or {{BODY}}), else appended.
    """
    out = output_path(path, overwrite, ".docx")
    fields = {k.strip("{} ").upper(): str(v) for k, v in (fields or {}).items()}
    if template:
        tpath = _find_template(template)
        d = docx.Document(str(tpath))
    else:
        d = docx.Document()
        tpath = None

    anchor = None
    for p in list(_all_paragraphs(d)):
        joined = "".join(r.text for r in p.runs)
        if re.search(r"\{\{\s*(INHALT|BODY|CONTENU)\s*\}\}", joined):
            anchor = p
            continue
        if "{{" in joined:
            _replace_in_paragraph(p, fields)

    if anchor is not None:
        for b in blocks:
            _add_block(d, b, before=anchor)
        anchor._element.getparent().remove(anchor._element)
    else:
        for b in blocks:
            _add_block(d, b)

    missing = sorted({m for p in _all_paragraphs(d) for m in re.findall(r"\{\{\s*([A-Z_ÄÖÜ0-9]+)\s*\}\}", p.text)})
    d.save(out)
    log_event("files", action="create", path=rel(out), tool="create_docx", template=str(tpath) if tpath else None)
    res: dict[str, Any] = {"path": rel(out)}
    if missing:
        res["unfilled_placeholders"] = missing
    return res


def _all_paragraphs(d):
    yield from d.paragraphs
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for s in d.sections:
        for part in (s.header, s.footer, s.first_page_header, s.first_page_footer):
            if part is not None and not part.is_linked_to_previous:
                yield from part.paragraphs
                for t in part.tables:
                    for row in t.rows:
                        for cell in row.cells:
                            yield from cell.paragraphs


def _replace_in_paragraph(p, fields: dict[str, str]) -> None:
    runs = p.runs
    if not runs:
        return
    joined = "".join(r.text for r in runs)

    def sub(m):
        key = m.group(1).strip().upper()
        return fields.get(key, m.group(0))

    new = re.sub(r"\{\{\s*([^}]+?)\s*\}\}", sub, joined)
    if new == joined:
        return
    lines = new.split("\n")
    runs[0].text = lines[0]
    for r in runs[1:]:
        r.text = ""
    for line in lines[1:]:
        runs[0].add_break()
        runs[0].add_text(line)


def _add_block(d, b: dict[str, Any], before=None) -> None:
    kind = (b.get("type") or "paragraph").lower()

    def new_par(style=None):
        if before is not None:
            return before.insert_paragraph_before(style=style) if _has_style(d, style) else before.insert_paragraph_before()
        return d.add_paragraph(style=style) if _has_style(d, style) else d.add_paragraph()

    if kind == "heading":
        level = int(b.get("level", 1))
        p = new_par(f"Heading {min(max(level, 1), 4)}")
        _add_runs(p, b.get("text", ""))
        if not _has_style(d, f"Heading {level}"):
            for r in p.runs:
                r.bold = True
    elif kind in ("bullets", "numbered"):
        style = "List Bullet" if kind == "bullets" else "List Number"
        for i, item in enumerate(b.get("items") or []):
            p = new_par(style)
            if not _has_style(d, style):
                item = ("• " if kind == "bullets" else f"{i + 1}. ") + str(item)
            _add_runs(p, str(item))
    elif kind == "table":
        rows = b.get("rows") or []
        if not rows:
            return
        ncols = max(len(r) for r in rows)
        table = d.add_table(rows=len(rows), cols=ncols)
        if _has_style(d, "Table Grid"):
            table.style = "Table Grid"
        for i, row in enumerate(rows):
            for j in range(ncols):
                cell = table.cell(i, j)
                cell.text = "" if j >= len(row) or row[j] is None else str(row[j])
                if i == 0 and b.get("header", True):
                    for r in cell.paragraphs[0].runs:
                        r.bold = True
        if before is not None:
            before._element.addprevious(table._element)
    elif kind == "page_break":
        p = new_par()
        p.add_run().add_break(WD_BREAK.PAGE)
    else:
        text = str(b.get("text", ""))
        for chunk in text.split("\n\n"):
            p = new_par()
            _add_runs(p, chunk)


def _add_runs(p, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            p.add_run(part[2:-2]).bold = True
        else:
            lines = part.split("\n")
            run = p.add_run(lines[0])
            for line in lines[1:]:
                run.add_break()
                run.add_text(line)


def _has_style(d, name) -> bool:
    if not name:
        return False
    try:
        d.styles[name]
        return True
    except KeyError:
        return False


@server.tool()
def create_pdf(path: str, blocks: list[dict[str, Any]], title: str | None = None, overwrite: bool = False) -> dict[str, Any]:
    """Create a PDF directly from blocks (same block format as create_docx: heading, paragraph, bullets,
    numbered, table, page_break)."""
    out = output_path(path, overwrite, ".pdf")
    _render_pdf(out, blocks, title)
    log_event("files", action="create", path=rel(out), tool="create_pdf")
    return {"path": rel(out)}


@server.tool()
def convert_to_pdf(path: str, overwrite: bool = False) -> dict[str, Any]:
    """Convert a .docx (or .xlsx) in the workspace to PDF. Uses LibreOffice if installed for an exact
    rendering; otherwise renders the document text and tables with a simple layout."""
    src = resolve_in_workspace(path, must_exist=True)
    out = output_path(str(src.with_suffix(".pdf")), overwrite, ".pdf")
    soffice = shutil.which("soffice") or _mac_soffice()
    if soffice:
        with tempfile.TemporaryDirectory() as tmp:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(src)],
                check=True, capture_output=True, timeout=120,
            )
            produced = Path(tmp) / (src.stem + ".pdf")
            shutil.move(str(produced), out)
        method = "libreoffice"
    else:
        if src.suffix.lower() != ".docx":
            raise ValueError("Without LibreOffice only .docx can be converted")
        _render_pdf(out, _docx_to_blocks(src), None)
        method = "simple-layout"
    log_event("files", action="create", path=rel(out), tool="convert_to_pdf", method=method)
    return {"path": rel(out), "method": method}


def _mac_soffice() -> str | None:
    p = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    return p if os.path.exists(p) else None


def _docx_to_blocks(src: Path) -> list[dict[str, Any]]:
    d = docx.Document(str(src))
    blocks: list[dict[str, Any]] = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = DocxParagraph(child, d)
            style = (p.style.name if p.style is not None else "") or ""
            if style.startswith("Heading"):
                lvl = int(style.split()[-1]) if style.split()[-1].isdigit() else 1
                blocks.append({"type": "heading", "text": p.text, "level": lvl})
            elif style.startswith("List"):
                blocks.append({"type": "bullets", "items": [p.text]})
            else:
                blocks.append({"type": "paragraph", "text": p.text})
        elif tag == "tbl":
            t = DocxTable(child, d)
            blocks.append({"type": "table", "rows": [[c.text for c in r.cells] for r in t.rows]})
    return blocks


def _render_pdf(out: Path, blocks: list[dict[str, Any]], title: str | None) -> None:
    styles = getSampleStyleSheet()

    def md(text: str) -> str:
        t = escape(str(text))
        t = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", t)
        return t.replace("\n", "<br/>")

    story: list = []
    if title:
        story += [PdfParagraph(md(title), styles["Title"]), Spacer(1, 4 * mm)]
    for b in blocks:
        kind = (b.get("type") or "paragraph").lower()
        if kind == "heading":
            lvl = min(max(int(b.get("level", 1)), 1), 3)
            story.append(PdfParagraph(md(b.get("text", "")), styles[f"Heading{lvl}"]))
        elif kind in ("bullets", "numbered"):
            items = [ListItem(PdfParagraph(md(i), styles["BodyText"])) for i in b.get("items") or []]
            story.append(ListFlowable(items, bulletType="bullet" if kind == "bullets" else "1"))
        elif kind == "table":
            rows = [[PdfParagraph(md("" if c is None else c), styles["BodyText"]) for c in r] for r in b.get("rows") or []]
            if rows:
                t = PdfTable(rows, repeatRows=1 if b.get("header", True) else 0)
                st = [("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]
                if b.get("header", True):
                    st.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6ECF2")))
                t.setStyle(TableStyle(st))
                story += [t, Spacer(1, 3 * mm)]
        elif kind == "page_break":
            story.append(PageBreak())
        else:
            for chunk in str(b.get("text", "")).split("\n\n"):
                story.append(PdfParagraph(md(chunk), styles["BodyText"]))
    doc = SimpleDocTemplate(str(out), pagesize=A4, leftMargin=22 * mm, rightMargin=20 * mm, topMargin=20 * mm, bottomMargin=20 * mm)
    doc.build(story or [PdfParagraph("", styles["BodyText"])])


@server.tool()
def move_file(source: str, destination: str) -> dict[str, Any]:
    """Move or rename a file inside the workspace (creates missing folders). Never overwrites: if the
    destination exists a _v2 suffix is added. Every move is logged to .gotthard/logs/moves.jsonl so it can be undone."""
    src = resolve_in_workspace(source, must_exist=True)
    if src.is_dir():
        raise WorkspaceError("Only files can be moved with this tool")
    dst = resolve_in_workspace(destination)
    if destination.endswith(("/", os.sep)) or dst.is_dir():
        dst = dst / src.name
    if {".gotthard", ".git"} & set(dst.relative_to(workspace_root()).parts):
        raise WorkspaceError("Moving into .gotthard/ or .git/ is not allowed")
    dst.parent.mkdir(parents=True, exist_ok=True)
    final = versioned_path(dst)
    shutil.move(str(src), str(final))
    log_event("moves", source=rel(src), destination=rel(final))
    return {"from": rel(src), "to": rel(final), "renamed_to_avoid_overwrite": final != dst}


def _template_dirs() -> list[Path]:
    dirs = []
    root = workspace_root()
    for name in ("Vorlagen", "Modeles", "Modèles", "Templates"):
        if (root / name).is_dir():
            dirs.append(root / name)
    for d in (os.environ.get(TEMPLATE_DIRS_ENV) or "").split(os.pathsep):
        if d and Path(d).is_dir():
            dirs.append(Path(d))
    return dirs


def _find_template(name: str) -> Path:
    cand = Path(name)
    if cand.is_absolute() and cand.exists():
        if any(d.resolve() in cand.resolve().parents for d in _template_dirs()):
            return cand
        raise WorkspaceError("Template must be in the workspace or a bundled template folder")
    try:
        p = resolve_in_workspace(name)
        if p.exists():
            return p
    except WorkspaceError:
        pass
    for d in _template_dirs():
        if (d / cand.name).exists():
            return d / cand.name
    available = sorted({f.name for d in _template_dirs() for f in d.glob("*.docx") if not f.name.startswith("~$")})
    raise WorkspaceError(f"Template not found: {name}. Available: {', '.join(available) or 'none'}")


def main() -> None:
    root = workspace_root()
    if not root.is_dir():
        raise SystemExit(f"Workspace does not exist: {root}")
    server.run("stdio")


if __name__ == "__main__":
    main()
