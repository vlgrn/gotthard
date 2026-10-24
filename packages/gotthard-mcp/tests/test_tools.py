from pathlib import Path

import openpyxl
import pytest

from gotthard_mcp import memory, office, readers
from gotthard_mcp.common import WorkspaceError, resolve_in_workspace

TEMPLATES = Path(__file__).resolve().parents[2] / "gotthard-skills" / "templates"


@pytest.fixture()
def ws(tmp_path, monkeypatch):
    monkeypatch.setenv("GOTTHARD_WORKSPACE", str(tmp_path))
    monkeypatch.setenv("GOTTHARD_TEMPLATES", str(TEMPLATES))
    monkeypatch.delenv("GOTTHARD_EMBED_BASE_URL", raising=False)
    return tmp_path


def test_confinement(ws):
    assert resolve_in_workspace("a/b.txt") == ws / "a" / "b.txt"
    for bad in ("../x.txt", "/etc/passwd", "a/../../x"):
        with pytest.raises(WorkspaceError):
            resolve_in_workspace(bad)


def test_xlsx_versioning_and_formats(ws):
    sheet = {"name": "Konten", "columns": ["Datum", "Text", "Betrag"],
             "rows": [["03.01.2026", "Miete", -2850.0], ["2026-01-08", "Gutschrift", "4'320.00"]], "totals": ["C"]}
    r1 = office.create_xlsx("out/Summary.xlsx", [sheet])
    r2 = office.create_xlsx("out/Summary.xlsx", [sheet])
    assert r1["path"] == "out/Summary.xlsx" and r2["path"] == "out/Summary_v2.xlsx"
    ws_ = openpyxl.load_workbook(ws / "out/Summary.xlsx")["Konten"]
    assert ws_["A2"].number_format == "DD.MM.YYYY"
    assert ws_["C3"].value == 4320.0
    assert ws_["C4"].value == "=SUM(C2:C3)"

    office.create_xlsx("serial.xlsx", [{"columns": ["Datum", "Betrag"], "rows": [[46025, 1.0]], "number_format": {"A": "@"}}])
    c = openpyxl.load_workbook(ws / "serial.xlsx").active["A2"]
    assert c.value.strftime("%d.%m.%Y") == "03.01.2026" and c.number_format == "DD.MM.YYYY"


def test_docx_template_fill(ws):
    res = office.create_docx(
        "Brief.docx",
        blocks=[{"type": "paragraph", "text": "Es fehlen **drei** Belege."}, {"type": "bullets", "items": ["A", "B"]}],
        template="Briefvorlage.docx",
        fields={"FIRMA": "Treuhand Beispiel GmbH", "BETREFF": "MWST Q4", "EMPFAENGER": "Müller AG\nBahnhofstrasse 12"},
    )
    text = readers.to_text(readers.extract(ws / res["path"]))
    assert "MWST Q4" in text and "drei" in text and "{{INHALT}}" not in text
    assert "ORT_DATUM" in res["unfilled_placeholders"]


def test_pdf_and_convert(ws):
    office.create_docx("x.docx", blocks=[{"type": "heading", "text": "Titel"}, {"type": "table", "rows": [["a", "b"], ["1", "2"]]}])
    assert (ws / office.create_pdf("r.pdf", [{"type": "paragraph", "text": "Hallo"}])["path"]).exists()
    assert (ws / office.convert_to_pdf("x.docx")["path"]).exists()


def test_move_never_overwrites(ws):
    (ws / "Eingang").mkdir()
    (ws / "Eingang/a.pdf").write_text("1")
    (ws / "Eingang/b.pdf").write_text("2")
    office.move_file("Eingang/a.pdf", "Klienten/X/2026-01-01_Rechnung.pdf")
    r = office.move_file("Eingang/b.pdf", "Klienten/X/2026-01-01_Rechnung.pdf")
    assert r["to"] == "Klienten/X/2026-01-01_Rechnung_v2.pdf"
    assert (ws / ".gotthard/logs/moves.jsonl").read_text().count("\n") == 2


def test_memory_remember_forget_and_search(ws):
    (ws / ".gotthard").mkdir(exist_ok=True)
    (ws / ".gotthard/memory.md").write_text("# Firmengedächtnis\n\n## Firma\nName: Test\n\n## Notizen\n")
    memory.remember("Pauschalhonorar CHF 2'400 für Abschluss 2025", subject="Keller GmbH")
    memory.remember("Bevorzugt Briefpost", subject="Müller AG")
    md = (ws / ".gotthard/memory.md").read_text()
    assert "Keller GmbH" in md and md.index("## Notizen") < md.index("Keller GmbH")

    (ws / "Keller").mkdir()
    (ws / "Keller/mail.txt").write_text("Honorar Jahresabschluss vereinbart")
    memory.indexer.index()
    assert memory.indexer.search("Jahresabschluss Honorar")[0]["path"] == "Keller/mail.txt"

    res = memory.forget("Keller")
    assert res["notes_removed"] == 1
    assert "Keller GmbH" not in (ws / ".gotthard/memory.md").read_text()
    assert "Müller AG" in (ws / ".gotthard/memory.md").read_text()
    memory.indexer.index()
    assert memory.indexer.search("Jahresabschluss Honorar") == []


def test_read_bundled_template_but_not_arbitrary_files(ws):
    assert "EMPFAENGER" in office.read_document("Briefvorlage.docx")
    with pytest.raises(WorkspaceError):
        office.read_document("/etc/hosts")
